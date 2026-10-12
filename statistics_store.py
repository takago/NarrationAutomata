"""Project-scoped append-only request statistics and on-demand XLSX export.

The recorder uses contextvars so concurrent NiceGUI worker threads cannot mix projects.
The XLSX exporter imports openpyxl lazily; recording does not require it.
"""
from __future__ import annotations

from collections import defaultdict
from contextlib import contextmanager
from contextvars import ContextVar
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import threading
import uuid

_stats_path: ContextVar[str | None] = ContextVar('slide_stats_path', default=None)
_page: ContextVar[int | None] = ContextVar('slide_stats_page', default=None)
_run_id: ContextVar[str | None] = ContextVar('slide_stats_run_id', default=None)
_write_lock = threading.Lock()


@contextmanager
def statistics_context(path: Path | str, page: int | None = None):
    tok_path = _stats_path.set(str(path))
    tok_page = _page.set(page)
    tok_run = _run_id.set(uuid.uuid4().hex)
    try:
        yield
    finally:
        _run_id.reset(tok_run)
        _page.reset(tok_page)
        _stats_path.reset(tok_path)


def run_with_statistics(path, statistics_page, function, *args, **kwargs):
    """Record statistics for a page without consuming the wrapped function's `page` keyword."""
    with statistics_context(path, statistics_page):
        return function(*args, **kwargs)


def set_current_page(page: int | None):
    _page.set(page)


def record(task: str, model: str, status: str, *, page=None, **details):
    location = _stats_path.get() or os.environ.get('SLIDE_STATISTICS_PATH')
    if not location:
        return
    path = Path(location)
    record = {
        'schema_version': 1,
        'timestamp': datetime.now(timezone.utc).astimezone().isoformat(timespec='seconds'),
        'run_id': _run_id.get() or os.environ.get('SLIDE_STATISTICS_RUN_ID') or 'cli',
        'request_id': uuid.uuid4().hex,
        'project_id': path.parent.name,
        'page': page if page is not None else _page.get(),
        'task': task, 'model': model, 'status': status,
        **details,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = (json.dumps(record, ensure_ascii=False, separators=(',', ':'), default=str) + '\n').encode('utf-8')
    # A single os.write on an O_APPEND descriptor. Lock protects threads in this
    # process; flock also protects concurrent CLI/UI processes on Linux.
    with _write_lock:
        fd = os.open(str(path), os.O_CREAT | os.O_APPEND | os.O_WRONLY, 0o600)
        try:
            try:
                import fcntl
                fcntl.flock(fd, fcntl.LOCK_EX)
            except ImportError:
                pass
            written = os.write(fd, payload)
            if written != len(payload):
                raise OSError('statistics.jsonlへの書き込みが不完全です')
        finally:
            os.close(fd)


def read_records(path: Path):
    if not path.exists():
        return []
    rows = []
    with path.open(encoding='utf-8') as file:
        for no, line in enumerate(file, 1):
            if not line.strip():
                continue
            try:
                obj = json.loads(line)
            except json.JSONDecodeError:
                # A torn final record is skipped; corrupt earlier entries are errors.
                if no == sum(1 for _ in path.open(encoding='utf-8')):
                    break
                raise ValueError(f'統計ファイルの{no}行目が破損しています')
            if isinstance(obj, dict):
                rows.append(obj)
    return rows


def export_xlsx(source: Path, destination: Path) -> Path:
    """Generate a readable multi-sheet Excel report from accumulated JSONL statistics."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    records = read_records(source)
    if not records:
        raise ValueError('統計情報はまだ記録されていません')
    wb = Workbook()
    wb.remove(wb.active)
    cols = ['timestamp','run_id','request_id','project_id','page','task','model','status',
            'reasoning_effort','paragraph_index','paragraph_total','chunk_index','chunk_total',
            'reasoning_chars','content_chars','total_chars','prompt_tokens','completion_tokens',
            'total_tokens','elapsed_sec','input_chars','audio_bytes','audio_duration_sec','rtf','error']

    def flat(record):
        values = dict(record)
        for group in ('llm','tts','filter'):
            subsection = values.pop(group, None)
            if isinstance(subsection, dict):
                values.update(subsection)
        return values

    rows = [flat(r) for r in records]

    def clean(value):
        if isinstance(value, (int, float, str, bool)) or value is None:
            return value
        return json.dumps(value, ensure_ascii=False, default=str)

    def sheet(name, headers, data):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in data:
            ws.append([clean(row.get(key)) for key in headers])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='17324D')
            cell.font = Font(bold=True, color='FFFFFF')
            cell.alignment = Alignment(vertical='center')
        ws.row_dimensions[1].height = 24
        for ix, key in enumerate(headers, 1):
            width = 44 if key == 'error' else 26 if key in ('timestamp', 'task', 'model', 'project_id', 'run_id', 'request_id') else 18
            ws.column_dimensions[get_column_letter(ix)].width = width
        return ws

    sheet('実行履歴', cols, rows)
    groups = defaultdict(list)
    for row in rows:
        for field in ('page', 'task', 'model'):
            if row.get(field) is not None:
                groups[(field, str(row[field]))].append(row)
    for field, title in [('page', 'スライド別'), ('task', '処理別'), ('model', 'モデル別')]:
        data = []
        for (key, value), values in sorted(groups.items(), key=lambda kv: kv[0][1]):
            if key != field:
                continue
            durations = [r.get('elapsed_sec') for r in values if isinstance(r.get('elapsed_sec'), (int, float))]
            data.append({'分類': int(value) if field == 'page' and value.isdecimal() else value,
                         '実行回数': len(values), '成功数': sum(r.get('status') == 'success' for r in values),
                         '失敗数': sum(r.get('status') != 'success' for r in values),
                         '合計時間(秒)': round(sum(durations), 2),
                         '平均時間(秒)': round(sum(durations)/len(durations), 2) if durations else None,
                         '思考文字数': sum(r.get('reasoning_chars') or 0 for r in values),
                         '回答文字数': sum(r.get('content_chars') or 0 for r in values)})
        sheet(title, ['分類','実行回数','成功数','失敗数','合計時間(秒)','平均時間(秒)','思考文字数','回答文字数'], data)
    sheet('TTS分析', cols, [r for r in rows if r.get('task') in ('tts', 'tts_filter')])
    sheet('エラー履歴', cols, [r for r in rows if r.get('status') != 'success'])
    destination.parent.mkdir(parents=True, exist_ok=True)
    temp = destination.with_name('.' + destination.name)
    try:
        wb.save(temp)
        temp.replace(destination)
    finally:
        temp.unlink(missing_ok=True)
    return destination
